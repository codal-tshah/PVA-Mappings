import zipfile
import tempfile
import shutil
from pathlib import Path

from pyxlsb import open_workbook
from openpyxl import Workbook


def find_xlsb(path):
    path = Path(path)

    if path.suffix.lower() == ".xlsb":
        return str(path), None

    if path.suffix.lower() != ".zip":
        raise ValueError("Input must be .xlsb or .zip")

    temp_dir = tempfile.mkdtemp(prefix="xlsb_extract_")

    with zipfile.ZipFile(path, "r") as z:
        z.extractall(temp_dir)

    xlsb_files = list(Path(temp_dir).rglob("*.xlsb"))

    if not xlsb_files:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise FileNotFoundError("No .xlsb file found in ZIP")

    return str(xlsb_files[0]), temp_dir


def convert_xlsb_to_xlsx(input_file, output_file=None):
    xlsb_path, temp_dir = find_xlsb(input_file)

    if output_file is None:
        output_file = str(Path(xlsb_path).with_suffix(".xlsx"))

    wb_out = Workbook(write_only=True)

    with open_workbook(xlsb_path) as wb_in:

        total_sheets = len(wb_in.sheets)

        for idx, sheet_name in enumerate(wb_in.sheets, start=1):

            print(f"[{idx}/{total_sheets}] Processing: {sheet_name}")

            ws_out = wb_out.create_sheet(title=sheet_name[:31])

            with wb_in.get_sheet(sheet_name) as ws_in:

                for row in ws_in.rows():
                    ws_out.append(
                        [cell.v for cell in row]
                    )

    if "Sheet" in wb_out.sheetnames:
        del wb_out["Sheet"]

    wb_out.save(output_file)

    if temp_dir:
        shutil.rmtree(temp_dir, ignore_errors=True)

    print(f"\nSaved: {output_file}")


if __name__ == "__main__":
    convert_xlsb_to_xlsx(
        "/Users/tshah/Documents/PVA Mappings/documents/PVA_Workbench.xlsb",
        "workbook.xlsx"
    )