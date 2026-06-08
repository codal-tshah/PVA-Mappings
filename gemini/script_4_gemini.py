import csv
import logging
import os
import re
from collections import defaultdict

import openpyxl
from openpyxl.utils import get_column_letter

# --- CONFIGURATION ---
FILE_PATH = "Mixed Use Copy zip real.xlsm"  # Ensure this points to your file
OUTPUT_DIR = "Workbench_Analysis_4_gemini"

TARGET_TABS = [
    "File Info", "Settings", "Dates, Premises", "Contracts, History", "Scope",
    "Site", "Zoning", "Improvements", "Assessment", "Land Grid",
    "Land Valuation", "Land Comp Profiles", "Cost Approach Setup", "Sales Grid",
    "Sales Approach", "Sale Comp Profiles", "Income Approach Setup", "Rent Roll Input",
    "Rent Roll Config", "Lease Grid", "Market Rent", "MF Market Rent",
    "Rent Comp Profiles", "Capitalization and Multipliers",
    "Investor Survey Input", "Investor Survey Output", "MF Lease Up",
    "DirectCapConclusion"
]

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")


class WorkbenchAnalyzer:
    def __init__(self, filepath, target_tabs, out_dir):
        self.filepath = filepath
        self.target_tabs = target_tabs
        self.out_dir = out_dir
        self.wb = None

        # Map "Sheet!A1" -> named range name
        self.named_range_map = {}

        # Track relationships between tabs
        self.relationships = set()

        if not os.path.exists(out_dir):
            os.makedirs(out_dir)

        self._init_csv_headers()

    def _init_csv_headers(self):
        headers = {
            "Tab Inventory.csv": ["Tab Name", "Category", "Purpose", "Input / Calc / Output", "Notes", "Visibility", "Tab Color"],
            "Field Inventory.csv": ["Tab Name", "Field Name", "Input", "Calculated", "Dropdown", "Data Type", "Notes", "Cell Ref"],
            "Field Mapping.csv": ["Source Tab", "Source Field", "Used By Tab", "Used For"],
            "Dropdown Values.csv": ["Tab Name", "Field Name", "Dropdown Values", "Cell Ref"],
            "Calculated Fields.csv": ["Tab Name", "Field Name", "Formula Logic", "Depends On", "Cell Ref"],
            "Tab Relationships.csv": ["Source Tab", "Target Tab", "Relationship"],
            "Notes Findings.csv": ["Tab Name", "Finding", "Open Question"],
            "ShowHide Usage.csv": ["Tab Name", "Cell Ref", "Formula", "Notes"]
        }
        for filename, cols in headers.items():
            filepath = os.path.join(self.out_dir, filename)
            if not os.path.exists(filepath):
                with open(filepath, 'w', newline='', encoding='utf-8') as f:
                    csv.writer(f).writerow(cols)

    def _append_to_csv(self, filename, row_data):
        with open(os.path.join(self.out_dir, filename), 'a', newline='', encoding='utf-8') as f:
            csv.writer(f).writerow([str(item) if item is not None else "" for item in row_data])

    def load_workbook(self):
        logging.info(f"Loading workbook {self.filepath} (This may take several minutes)...")
        self.wb = openpyxl.load_workbook(self.filepath, data_only=False, keep_vba=True)
        logging.info("Workbook loaded successfully.")
        self._build_named_range_map()

    def _build_named_range_map(self):
        """Maps named ranges to cell coordinates (e.g., 'File Info!I13' -> 'Z_FileInfo_Date')"""
        logging.info("Building Named Range index...")
        try:
            defined_names = self.wb.defined_names.items()
        except AttributeError:
            defined_names = [(dn.name, dn) for dn in getattr(self.wb.defined_names, 'definedName', [])]

        for name, dn in defined_names:
            try:
                if getattr(dn, 'type', 'REGION') != "REGION" or getattr(dn, 'is_external', False):
                    continue
                
                val_str = str(dn.value)
                for part in val_str.split(','):
                    match = re.search(r"'?([^']+)'?!\$?([A-Za-z]+)\$?([0-9]+)", part.strip())
                    if match:
                        sheet_name, col_letter, row_num = match.groups()
                        key = f"{sheet_name}!{col_letter}{row_num}"
                        self.named_range_map[key] = name
            except Exception:
                continue

    def resolve_dropdown_values(self, formula1):
        """
        Attempts to resolve Data Validation references into actual comma-separated values.
        Handles Hardcoded strings ("A,B,C") and Named Ranges (=_Property_Types).
        """
        if not formula1:
            return ""

        # Remove leading equals sign
        ref = str(formula1).lstrip('=')

        # 1. HARDCODED LIST: It already has commas or quotes (e.g., '"None,Fannie,Freddie"')
        if ',' in ref or '"' in ref:
            return ref.replace('"', '')

        # 2. DYNAMIC CASCADING DROPDOWN (e.g., INDIRECT($I$92))
        if "INDIRECT" in ref.upper():
            return f"[Dynamic Formula: {formula1}]"

        # 3. NAMED RANGE RESOLUTION (e.g., _Property_Types)
        if ref in self.wb.defined_names:
            try:
                dn = self.wb.defined_names[ref]
                resolved_values = []
                
                # Fetch the destination coordinates of the Named Range
                for sheet_title, coord in dn.destinations:
                    if sheet_title in self.wb.sheetnames:
                        ws = self.wb[sheet_title]
                        
                        # If it's a range (A1:A10)
                        if ':' in coord:
                            for row in ws[coord]:
                                for cell in row:
                                    if cell.value is not None:
                                        resolved_values.append(str(cell.value).strip())
                        # If it's a single cell
                        else:
                            val = ws[coord].value
                            if val is not None:
                                resolved_values.append(str(val).strip())
                                
                if resolved_values:
                    return ", ".join(resolved_values)
                else:
                    return f"[{ref}] (Empty Range)"
            except Exception:
                return f"[{ref}] (Error Resolving)"

        # Fallback if it's a complex formula we can't parse
        return formula1

    def find_field_label(self, ws, row, col):
        """Heuristic: Looks for a named range first, then scans left/up for a string label."""
        coord_key = f"{ws.title}!{get_column_letter(col)}{row}"
        if coord_key in self.named_range_map:
            return f"[{self.named_range_map[coord_key]}]"

        # Look Left (Search up to 7 columns away to handle merged/spaced labels)
        for c in range(col - 1, max(0, col - 8), -1):
            val = ws.cell(row=row, column=c).value
            if val and isinstance(val, str) and not val.startswith("="):
                clean_val = val.strip(" :>-")
                if len(clean_val) > 1: return clean_val
                
        # Look Above (Search up to 7 rows away)
        for r in range(row - 1, max(0, row - 8), -1):
            val = ws.cell(row=r, column=col).value
            if val and isinstance(val, str) and not val.startswith("="):
                clean_val = val.strip(" :>-")
                if len(clean_val) > 1: return clean_val
                
        return f"Unknown_Field_{get_column_letter(col)}{row}"

    def extract_dependencies_from_formula(self, formula):
        pattern = r"'([^']+)'!\$?[A-Za-z]+\$?[0-9]+|([A-Za-z0-9_]+)!\$?[A-Za-z]+\$?[0-9]+"
        matches = re.findall(pattern, formula)
        deps = set()
        for match in matches:
            dep = match[0] if match[0] else match[1]
            deps.add(dep)
        return list(deps)

    def is_pass_through(self, formula):
        """Identifies formulas that just pass data, e.g., =A1 or ='Sheet Name'!$A$1"""
        pattern = r"^='?[^']+'?!\$?[A-Za-z]+\$?[0-9]+$|^=\$?[A-Za-z]+\$?[0-9]+$"
        return bool(re.match(pattern, str(formula).strip()))

    def is_show_hide_formula(self, formula):
        """Catches the specific backend toggle formulas."""
        return "ShowHide" in str(formula) or "1/(" in str(formula)

    def analyze(self):
        self.load_workbook()
        
        for sheet_name in self.target_tabs:
            if sheet_name not in self.wb.sheetnames:
                logging.warning(f"Tab '{sheet_name}' not found. Skipping.")
                continue
                
            logging.info(f"Analyzing sheet: {sheet_name}")
            ws = self.wb[sheet_name]
            self.analyze_sheet(ws)
            logging.info(f"Finished saving data for: {sheet_name}")
            
        logging.info("Analysis Complete! Generating final Tab Relationships...")
        self._write_relationships()

    def analyze_sheet(self, ws):
        sheet_name = ws.title
        visibility = ws.sheet_state
        tab_color = ws.sheet_properties.tabColor.rgb if ws.sheet_properties.tabColor else "None"
        
        self._append_to_csv("Tab Inventory.csv", [
            sheet_name, "Uncategorized", "Auto-extracted", "", "", visibility, tab_color
        ])
        
        # EXTRACT AND RESOLVE DROPDOWNS
        dropdowns = {}
        for dv in ws.data_validations.dataValidation:
            if dv.type == "list":
                # RESOLVE THE ACTUAL VALUES HERE
                resolved_values = self.resolve_dropdown_values(dv.formula1)
                for cell_range in dv.sqref.ranges:
                    for row in range(cell_range.min_row, cell_range.max_row + 1):
                        for col in range(cell_range.min_col, cell_range.max_col + 1):
                            coord = f"{get_column_letter(col)}{row}"
                            dropdowns[coord] = resolved_values

        max_r, max_c = ws.max_row, ws.max_column
        if max_r > 5000: max_r = 5000 
        
        for row in range(1, max_r + 1):
            row_dim = ws.row_dimensions.get(row)
            is_hidden_row = row_dim and row_dim.hidden
                
            for col in range(1, max_c + 1):
                col_letter = get_column_letter(col)
                col_dim = ws.column_dimensions.get(col_letter)
                is_hidden_col = col_dim and col_dim.hidden

                cell = ws.cell(row=row, column=col)
                val = cell.value
                coord = cell.coordinate
                
                if val is None: continue

                is_formula = isinstance(val, str) and str(val).startswith("=")
                is_dropdown = coord in dropdowns
                
                # Filter criteria: Input field, Formula, or Dropdown
                if is_formula or is_dropdown or (not isinstance(val, str) and not cell.protection.locked):
                    label = self.find_field_label(ws, row, col)
                    data_type = cell.data_type
                    
                    notes = []
                    if is_hidden_row or is_hidden_col:
                        notes.append("[HIDDEN CELL]")
                    if is_formula and self.is_pass_through(val):
                        notes.append("[PASS-THROUGH]")
                        
                    notes_str = " | ".join(notes)
                    
                    self._append_to_csv("Field Inventory.csv", [
                        sheet_name, label, 
                        "Yes" if not is_formula else "No", 
                        "Yes" if is_formula else "No", 
                        "Yes" if is_dropdown else "No", 
                        data_type, notes_str, coord
                    ])
                    
                    if is_dropdown:
                        self._append_to_csv("Dropdown Values.csv", [
                            sheet_name, label, dropdowns[coord], coord
                        ])
                        
                    if is_formula:
                        deps = self.extract_dependencies_from_formula(val)
                        
                        if self.is_show_hide_formula(val):
                            self._append_to_csv("ShowHide Usage.csv", [
                                sheet_name, coord, val, "Backend UI Toggle"
                            ])
                        else:
                            self._append_to_csv("Calculated Fields.csv", [
                                sheet_name, label, val, ", ".join(deps), coord
                            ])
                        
                        for dep in deps:
                            if dep != sheet_name:
                                self.relationships.add((dep, sheet_name))
                                self._append_to_csv("Field Mapping.csv", [
                                    dep, "Extracted from formula", sheet_name, f"Used in {label} ({coord})"
                                ])

    def _write_relationships(self):
        for source, target in sorted(self.relationships):
            self._append_to_csv("Tab Relationships.csv", [
                source, target, f"'{target}' pulls data from '{source}'"
            ])

if __name__ == "__main__":
    analyzer = WorkbenchAnalyzer(FILE_PATH, TARGET_TABS, OUTPUT_DIR)
    analyzer.analyze()