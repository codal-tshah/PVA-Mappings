import csv
import logging
import os
import re
from collections import defaultdict
from datetime import date, datetime

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

# --- CONFIGURATION ---
FILE_PATH = "/Users/tshah/Documents/PVA Mappings/gemini/Mixed Use Copy zip real.xlsm"  # Ensure this points to your file
OUTPUT_DIR = "Workbench_Analysis_gpt4"

TARGET_TABS = [
    "File Info", "Settings", "Dates, Premises", "Contracts, History", "Scope",
    "Site", "Zoning", "Improvements", "Assessment", "Land Grid",
    "Land Valuation", "Land Comp Profiles", "Cost Approach Setup", "Sales Grid",
    "Sales Approach", "Sale Comp Profiles", "Income Approach Setup", "Rent Roll Input",
    "Rent Roll Config", "Lease Grid", "Market Rent", "MF Market Rent",
    "Rent Comp Profiles", "Capitalization and Multipliers",
    "Investor Survey Input", "Investor Survey Output", "MF Lease Up",
    "DirectCapConclusion"
]

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")


class WorkbenchAnalyzer:
    def __init__(self, filepath, target_tabs, out_dir):
        self.filepath = filepath
        self.target_tabs = target_tabs
        self.out_dir = out_dir
        self.wb = None

        # Map "Sheet!A1" -> named range name
        self.named_range_map = {}
        self.named_range_cell_map = {}
        self.named_range_usage = defaultdict(lambda: {"tabs": set(), "csvs": set()})
        self.table_map = {}
        self.effective_value_cache = {}
        self.field_label_cache = {}
        self.merged_cell_lookup = {}

        # Track relationships between tabs and their types
        self.relationships = defaultdict(set)

        # Track how many named ranges each sheet has
        self.sheet_named_range_count = defaultdict(int)
        self.csv_handles = {}
        self.csv_writers = {}

        if os.path.exists(out_dir):
            self._clear_existing_csvs()
        else:
            os.makedirs(out_dir)

        self._init_csv_headers()

    def _clear_existing_csvs(self):
        """Delete old CSVs so reruns do not append duplicate output."""
        for filename in os.listdir(self.out_dir):
            if filename.lower().endswith(".csv"):
                try:
                    os.remove(os.path.join(self.out_dir, filename))
                except Exception as exc:
                    logging.warning(f"Could not remove {filename}: {exc}")

    def _init_csv_headers(self):
        """Initializes the CSV files with their headers."""
        headers = {
            "Tab Inventory.csv": [
                "Tab Name",
                "Category",
                "Purpose",
                "Input / Calc / Output",
                "Notes",
                "Visibility",
                "Tab Color"
            ],

            "Field Inventory.csv": [
                "Tab Name",
                "Field Name",
                "Named Range",
                "Input",
                "Calculated",
                "Dropdown",
                "Formula Type",
                "Data Type",
                "Required",
                "Notes",
                "Cell Ref"
            ],

            "Field Mapping.csv": [
                "Source Tab",
                "Source Field",
                "Used By Tab",
                "Used For"
            ],

            "Dropdown Values.csv": [
                "Tab Name",
                "Field Name",
                "All dropdown options",
                "Named ranges",
                "Dropdown Type",
                "Source Formula",
                "Referenced Cells",
                "Source Ranges",
                "Resolution Status",
                "Cell Ref"
            ],

            "Calculated Fields.csv": [
                "Tab Name",
                "Field Name",
                "Formula Logic",
                "Depends On (Source Fields)",
                "Named Range Dependencies",
                "Cell Ref"
            ],

            "Named Ranges.csv": [
                "Named Range",
                "Sheet",
                "Start Cell",
                "End Cell",
                "Cell Count",
                "Values Preview",
                "Used By Tabs",
                "Used In CSVs",
                "Reference"
            ],

            "ShowHide Usage.csv": [
                "Sheet",
                "Cell",
                "Formula",
                "Named Range"
            ],

            "Tab Relationships.csv": [
                "Source Tab",
                "Target Tab",
                "Relationship Type",
                "Relationship"
            ],

            "Notes Findings.csv": [
                "Tab Name",
                "Finding",
                "Open Question"
            ]
        }

        for filename, cols in headers.items():
            filepath = os.path.join(self.out_dir, filename)
            with open(filepath, "w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerow(cols)

        for filename in headers:
            filepath = os.path.join(self.out_dir, filename)
            handle = open(filepath, "a", newline="", encoding="utf-8")
            self.csv_handles[filename] = handle
            self.csv_writers[filename] = csv.writer(handle)

    def _append_to_csv(self, filename, row_data):
        writer = self.csv_writers.get(filename)
        if writer is None:
            filepath = os.path.join(self.out_dir, filename)
            with open(filepath, "a", newline="", encoding="utf-8") as f:
                csv.writer(f).writerow([str(item) if item is not None else "" for item in row_data])
            return

        writer.writerow([str(item) if item is not None else "" for item in row_data])

    def _close_csv_handles(self):
        for handle in self.csv_handles.values():
            try:
                handle.flush()
                handle.close()
            except Exception:
                pass
        self.csv_handles.clear()
        self.csv_writers.clear()

    def load_workbook(self):
        logging.info(f"Loading workbook {self.filepath} (This may take several minutes)...")
        self.wb = openpyxl.load_workbook(self.filepath, data_only=False, keep_vba=True)
        logging.info("Workbook loaded successfully.")
        self._build_merged_cell_lookup()
        self._build_named_range_map()
        self._build_table_map()

    def _build_merged_cell_lookup(self):
        """
        Cache merged-cell coordinates so effective value lookups are O(1).
        """
        self.merged_cell_lookup.clear()

        for ws in self.wb.worksheets:
            for merged in ws.merged_cells.ranges:
                top_left = ws.cell(merged.min_row, merged.min_col).value
                for row in range(merged.min_row, merged.max_row + 1):
                    for col in range(merged.min_col, merged.max_col + 1):
                        self.merged_cell_lookup[f"{ws.title}!{get_column_letter(col)}{row}"] = top_left

    def _iter_defined_names(self):
        """
        Returns a uniform iterable of (name, defined_name_object) across openpyxl versions.
        """
        try:
            return list(self.wb.defined_names.items())
        except AttributeError:
            return [
                (dn.name, dn)
                for dn in getattr(self.wb.defined_names, "definedName", [])
            ]

    def _iter_defined_name_destinations(self, dn):
        """
        Yields (sheet_name, cell_or_range) for a named range.
        Uses dn.destinations when available; otherwise falls back to parsing dn.value / dn.attr_text.
        """
        try:
            destinations = list(dn.destinations)
            if destinations:
                for sheet_name, coord in destinations:
                    yield sheet_name, coord
                return
        except Exception:
            pass

        raw = getattr(dn, "value", None) or getattr(dn, "attr_text", None) or str(dn)
        raw = str(raw)

        if any(token in raw.upper() for token in ["OFFSET(", "INDEX(", "INDIRECT(", "COUNTA(", "[", "!", ":"]):
            try:
                resolved_meta = self.resolve_dropdown_options(raw)
                source_ranges = str(resolved_meta.get("source_ranges", "")).strip()
                if source_ranges:
                    for part in [p.strip() for p in source_ranges.split(",") if p.strip()]:
                        match = re.fullmatch(r"(?:'([^']+)'|([^'!]+))!(.+)", part)
                        if match:
                            quoted_sheet, unquoted_sheet, coord = match.groups()
                            sheet_name = quoted_sheet if quoted_sheet else unquoted_sheet
                            if sheet_name and coord:
                                yield sheet_name, coord
                    return
            except Exception:
                pass

        # Handles:
        #   'Sheet Name'!$A$1
        #   Sheet1!$A$1:$C$3
        pattern = r"(?:'([^']+)'|([^'!]+))!\$?([A-Za-z]+)\$?([0-9]+)(?::\$?([A-Za-z]+)\$?([0-9]+))?"
        for match in re.finditer(pattern, raw):
            quoted_sheet, unquoted_sheet, c1, r1, c2, r2 = match.groups()
            sheet_name = quoted_sheet if quoted_sheet else unquoted_sheet
            cell_ref = f"{c1}{r1}"
            if c2 and r2:
                cell_ref = f"{cell_ref}:{c2}{r2}"
            yield sheet_name, cell_ref

    def _get_defined_name(self, name):
        """
        Return a defined name object by name across openpyxl versions.
        """
        try:
            return self.wb.defined_names[name]
        except Exception:
            for dn_name, dn in self._iter_defined_names():
                if dn_name == name:
                    return dn
        return None

    def _resolve_cell_range_values(self, ws, coord):
        """
        Resolve a single cell or range into a list of non-empty string values.
        """
        resolved_values = []

        try:
            if ":" in coord:
                for row in ws[coord]:
                    for cell in row:
                        if cell.value is not None:
                            resolved_values.append(str(cell.value).strip())
            else:
                value = ws[coord].value
                if value is not None:
                    resolved_values.append(str(value).strip())
        except Exception:
            return []

        return [value for value in resolved_values if value != ""]

    def _split_formula_args(self, arg_text):
        parts = []
        current = []
        depth = 0
        in_quotes = False

        for char in arg_text:
            if char == '"' and (not current or current[-1] != "\\"):
                in_quotes = not in_quotes
                current.append(char)
                continue

            if not in_quotes:
                if char == "(":
                    depth += 1
                elif char == ")" and depth > 0:
                    depth -= 1
                elif char == "," and depth == 0:
                    part = "".join(current).strip()
                    if part:
                        parts.append(part)
                    current = []
                    continue

            current.append(char)

        tail = "".join(current).strip()
        if tail:
            parts.append(tail)
        return parts

    def _extract_referenced_cells(self, formula_text):
        """
        Pull out direct cell references appearing in a formula string.
        """
        refs = set()
        text = str(formula_text)

        sheet_ref_pattern = r"(?:'([^']+)'|([A-Za-z0-9_ .\-]+))!\$?[A-Za-z]{1,3}\$?[0-9]+"
        for quoted_sheet, unquoted_sheet in re.findall(sheet_ref_pattern, text):
            sheet = quoted_sheet if quoted_sheet else unquoted_sheet
            if sheet:
                refs.add(sheet.strip())

        cell_pattern = r"(?<![A-Z0-9_!])\$?[A-Z]{1,3}\$?[0-9]+(?![A-Z0-9_])"
        for ref in re.findall(cell_pattern, text):
            refs.add(ref.replace("$", ""))

        return sorted(refs)

    def _get_sheet_and_coord_from_ref(self, ref, current_sheet=None):
        """
        Parse a direct sheet/cell/range reference and return the target sheet plus coord.
        """
        if not ref:
            return None, None

        text = str(ref).strip().lstrip("=")

        range_pattern = r"^(?:'([^']+)'|([^'!]+))!\$?([A-Za-z]{1,3})\$?([0-9]+)(?::\$?([A-Za-z]{1,3})\$?([0-9]+))?$"
        match = re.fullmatch(range_pattern, text)
        if match:
            quoted_sheet, unquoted_sheet, c1, r1, c2, r2 = match.groups()
            sheet_name = quoted_sheet if quoted_sheet else unquoted_sheet
            coord = f"{c1}{r1}"
            if c2 and r2:
                coord = f"{coord}:{c2}{r2}"
            return sheet_name.strip(), coord

        local_range_pattern = r"^\$?([A-Za-z]{1,3})\$?([0-9]+)(?::\$?([A-Za-z]{1,3})\$?([0-9]+))?$"
        match = re.fullmatch(local_range_pattern, text)
        if match and current_sheet:
            c1, r1, c2, r2 = match.groups()
            coord = f"{c1}{r1}"
            if c2 and r2:
                coord = f"{coord}:{c2}{r2}"
            return current_sheet, coord

        return None, None

    def _count_non_empty_cells_in_ref(self, ref, current_sheet=None):
        """
        Return COUNTA-like cell count for a direct reference if possible.
        """
        sheet_name, coord = self._get_sheet_and_coord_from_ref(ref, current_sheet=current_sheet)
        if not sheet_name or not coord or sheet_name not in self.wb.sheetnames:
            return None

        ws = self.wb[sheet_name]
        try:
            if ":" in coord:
                count = 0
                for row in ws[coord]:
                    for cell in row:
                        if cell.value is not None and str(cell.value).strip() != "":
                            count += 1
                return count
            value = ws[coord].value
            return 1 if value is not None and str(value).strip() != "" else 0
        except Exception:
            return None

    def _evaluate_numeric_expr(self, expr, current_sheet=None):
        """
        Resolve a small subset of numeric expressions used in validations.
        """
        text = str(expr).strip().lstrip("=")
        if re.fullmatch(r"[+-]?\d+", text):
            return int(text)

        match = re.fullmatch(r"COUNTA\((.+)\)", text, re.IGNORECASE)
        if match:
            return self._count_non_empty_cells_in_ref(match.group(1), current_sheet=current_sheet)

        return None

    def _resolve_index_expression(self, expr, current_sheet=None):
        """
        Resolve a simple INDEX(range, row[, col]) expression to a single-cell reference.
        """
        text = str(expr).strip().lstrip("=")
        match = re.fullmatch(r"INDEX\((.+)\)", text, re.IGNORECASE)
        if not match:
            return None

        args = self._split_formula_args(match.group(1))
        if len(args) < 2:
            return None

        range_ref = args[0]
        row_expr = args[1]
        col_expr = args[2] if len(args) > 2 else "1"

        sheet_name, coord = self._get_sheet_and_coord_from_ref(range_ref, current_sheet=current_sheet)
        if not sheet_name or sheet_name not in self.wb.sheetnames:
            return None

        row_idx = self._evaluate_numeric_expr(row_expr, current_sheet=current_sheet)
        col_idx = self._evaluate_numeric_expr(col_expr, current_sheet=current_sheet)
        if row_idx is None or col_idx is None:
            return None

        # Default to the top-left cell if a full range was supplied.
        if ":" in coord:
            min_col, min_row, _, _ = range_boundaries(coord)
            base_row = min_row
            base_col = min_col
            resolved_row = base_row + row_idx - 1
            resolved_col = base_col + col_idx - 1
        else:
            resolved_row = row_idx
            resolved_col = col_idx

        cell_ref = f"{get_column_letter(resolved_col)}{resolved_row}"
        value = self.wb[sheet_name][cell_ref].value
        if value is None or str(value).strip() == "":
            return None

        return {
            "dropdown_type": "Dynamic Dropdown",
            "options": str(value).strip(),
            "named_ranges": "",
            "source_formula": text,
            "referenced_cells": ", ".join(self._extract_referenced_cells(text)),
            "source_ranges": f"{sheet_name}!{cell_ref}",
            "resolved": True,
            "is_dynamic": True,
        }

    def _resolve_structured_reference(self, ref):
        """
        Resolve simple structured table references like Table_Properties[Type].
        """
        text = str(ref).strip().lstrip("=")
        match = re.fullmatch(r"([A-Za-z_][\w.]*)\[(.+)\]", text)
        if not match:
            return None

        table_name, body = match.groups()
        info = self.table_map.get(table_name)
        if not info:
            return None

        column_tokens = re.findall(r"\[([^\[\]]+)\]", body)
        column_name = ""
        for token in column_tokens:
            token = token.strip()
            if token and not token.startswith("#"):
                column_name = token
        if not column_name:
            return None

        sheet_name = info["sheet"]
        if sheet_name not in self.wb.sheetnames:
            return None

        ws = self.wb[sheet_name]
        min_col, min_row, max_col, max_row = range_boundaries(info["ref"])
        columns = info["columns"] or []
        target_offset = None
        for idx, col_name in enumerate(columns):
            if str(col_name).strip().lower() == column_name.lower():
                target_offset = idx
                break

        if target_offset is None:
            return None

        target_col = min_col + target_offset
        values = []
        for row in range(min_row + 1, max_row + 1):
            value = ws.cell(row=row, column=target_col).value
            if value is not None and str(value).strip() != "":
                values.append(str(value).strip())

        if not values:
            return None

        coord = f"{get_column_letter(target_col)}{min_row + 1}:{get_column_letter(target_col)}{max_row}"
        return {
            "dropdown_type": "Structured Reference",
            "options": ", ".join(values),
            "named_ranges": "",
            "source_formula": text,
            "referenced_cells": "",
            "source_ranges": f"{sheet_name}!{coord}",
            "resolved": True,
            "is_dynamic": False,
        }

    def _resolve_dynamic_target(self, ref, current_sheet=None):
        """
        Best-effort resolution for INDIRECT/OFFSET/INDEX-driven references.
        """
        text = str(ref).strip().lstrip("=")

        # First, try to resolve the expression as a nested reference.
        nested_sheet, nested_coord = self._get_sheet_and_coord_from_ref(text, current_sheet=current_sheet)
        if nested_sheet and nested_coord and nested_sheet in self.wb.sheetnames:
            values = self._resolve_cell_range_values(self.wb[nested_sheet], nested_coord)
            if values:
                return {
                    "dropdown_type": "Dynamic Dropdown",
                    "options": ", ".join(values),
                    "named_ranges": "",
                    "source_formula": text,
                    "referenced_cells": ", ".join(self._extract_referenced_cells(text)),
                    "source_ranges": f"{nested_sheet}!{nested_coord}",
                    "resolved": True,
                    "is_dynamic": True,
                }

        # If the dynamic target is a cell reference, read the cell and resolve what it points to.
        sheet_name, coord = self._get_sheet_and_coord_from_ref(text, current_sheet=current_sheet)
        if sheet_name and coord and sheet_name in self.wb.sheetnames:
            cell_value = self.wb[sheet_name][coord].value
            if cell_value is not None:
                nested = self.resolve_dropdown_options(cell_value, current_sheet=sheet_name)
                if nested.get("resolved"):
                    nested["dropdown_type"] = "Dynamic Dropdown"
                    nested["source_formula"] = text
                    nested["referenced_cells"] = ", ".join(sorted(set(
                        (nested.get("referenced_cells", "") + ", " + coord).replace(" ", "").split(",")
                    )))
                    nested["is_dynamic"] = True
                    return nested

        # Fallback: record lineage even if we cannot resolve.
        named_ranges = []
        for nr in set(self.named_range_map.values()):
            if re.search(rf"(?<![A-Z0-9_]){re.escape(str(nr))}(?![A-Z0-9_])", text, re.IGNORECASE):
                named_ranges.append(nr)

        return {
            "dropdown_type": "Dynamic Dropdown",
            "options": "",
            "named_ranges": ", ".join(sorted(set(named_ranges))),
            "source_formula": text,
            "referenced_cells": ", ".join(self._extract_referenced_cells(text)),
            "source_ranges": "",
            "resolved": False,
            "is_dynamic": True,
        }

    def resolve_dropdown_options(self, formula1, current_sheet=None):
        """
        Resolve data validation list formulas into dropdown metadata.

        Returns:
            dict with dropdown lineage and resolution metadata.
        """
        if not formula1:
            return {
                "dropdown_type": "Unresolved",
                "options": "",
                "named_ranges": "",
                "source_formula": "",
                "referenced_cells": "",
                "source_ranges": "",
                "resolved": False,
                "is_dynamic": False,
            }

        ref = str(formula1).strip().lstrip("=")
        if not ref:
            return {
                "dropdown_type": "Unresolved",
                "options": "",
                "named_ranges": "",
                "source_formula": "",
                "referenced_cells": "",
                "source_ranges": "",
                "resolved": False,
                "is_dynamic": False,
            }

        # Hardcoded list validation, e.g. "A,B,C"
        if "!" not in ref and ("," in ref or '"' in ref):
            return {
                "dropdown_type": "Static List",
                "options": ref.replace('"', ""),
                "named_ranges": "",
                "source_formula": ref,
                "referenced_cells": "",
                "source_ranges": "",
                "resolved": True,
                "is_dynamic": False,
            }

        if "INDIRECT(" in ref.upper():
            return self._resolve_dynamic_target(ref, current_sheet=current_sheet)

        # OFFSET(start, rows, cols, height, width)
        if ref.upper().startswith("OFFSET(") and ref.endswith(")"):
            inner = ref[7:-1]
            args = self._split_formula_args(inner)
            if len(args) >= 4:
                base_ref = args[0]
                row_offset = args[1]
                col_offset = args[2]
                height_expr = args[3]
                width_expr = args[4] if len(args) > 4 else "1"

                base_sheet, base_coord = self._get_sheet_and_coord_from_ref(base_ref, current_sheet=current_sheet)
                if base_sheet and base_coord and base_sheet in self.wb.sheetnames:
                    try:
                        if ":" in base_coord:
                            min_col, min_row, _, _ = range_boundaries(base_coord)
                            base_row = min_row
                            base_col = min_col
                        else:
                            base_col = column_index_from_string(re.match(r"[A-Za-z]{1,3}", base_coord).group(0))
                            base_row = int(re.search(r"\d+", base_coord).group(0))

                        row_offset_val = self._evaluate_numeric_expr(row_offset, current_sheet=current_sheet)
                        col_offset_val = self._evaluate_numeric_expr(col_offset, current_sheet=current_sheet)
                        height_val = self._evaluate_numeric_expr(height_expr, current_sheet=current_sheet)
                        width_val = self._evaluate_numeric_expr(width_expr, current_sheet=current_sheet)

                        if None not in (row_offset_val, col_offset_val, height_val, width_val):
                            start_row = base_row + row_offset_val
                            start_col = base_col + col_offset_val
                            end_row = start_row + height_val - 1
                            end_col = start_col + width_val - 1
                            coord = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"
                            values = self._resolve_cell_range_values(self.wb[base_sheet], coord)
                            if values:
                                return {
                                    "dropdown_type": "Dynamic Dropdown",
                                    "options": ", ".join(values),
                                    "named_ranges": "",
                                    "source_formula": ref,
                                    "referenced_cells": ", ".join(self._extract_referenced_cells(ref)),
                                    "source_ranges": f"{base_sheet}!{coord}",
                                    "resolved": True,
                                    "is_dynamic": True,
                                }
                    except Exception:
                        pass

            return self._resolve_dynamic_target(ref, current_sheet=current_sheet)

        # INDEX(range, n [, m]) used directly or inside a start:end expression.
        if ref.upper().startswith("INDEX("):
            # direct INDEX(range,row[,col]) -> resolve to a single cell reference when possible
            resolved = self._resolve_index_expression(ref, current_sheet=current_sheet)
            if resolved:
                return resolved

        if ":" in ref and "INDEX(" in ref.upper():
            start_ref, end_ref = ref.split(":", 1)
            end_resolved = self._resolve_index_expression(end_ref, current_sheet=current_sheet)
            start_sheet, start_coord = self._get_sheet_and_coord_from_ref(start_ref, current_sheet=current_sheet)
            if end_resolved and start_sheet and start_coord and start_sheet in self.wb.sheetnames:
                end_text = end_resolved["source_ranges"].split("!", 1)[-1]
                range_coord = f"{start_coord.split(':', 1)[0]}:{end_text}"
                values = self._resolve_cell_range_values(self.wb[start_sheet], range_coord)
                if values:
                    return {
                        "dropdown_type": "Dynamic Dropdown",
                        "options": ", ".join(values),
                        "named_ranges": "",
                        "source_formula": ref,
                        "referenced_cells": ", ".join(sorted(set(
                            self._extract_referenced_cells(ref) + end_resolved.get("referenced_cells", "").split(", ")
                        ))).strip(", "),
                        "source_ranges": f"{start_sheet}!{range_coord}",
                        "resolved": True,
                        "is_dynamic": True,
                    }

        # Named range validation.
        dn = self._get_defined_name(ref)
        if dn is not None:
            resolved_values = []
            source_ranges = []
            try:
                for sheet_name, coord in self._iter_defined_name_destinations(dn):
                    if sheet_name in self.wb.sheetnames:
                        source_ranges.append(f"{sheet_name}!{coord}")
                        resolved_values.extend(
                            self._resolve_cell_range_values(self.wb[sheet_name], coord)
                        )
            except Exception:
                resolved_values = []

            if resolved_values:
                return {
                    "dropdown_type": "Named Range",
                    "options": ", ".join(resolved_values),
                    "named_ranges": ref,
                    "source_formula": ref,
                    "referenced_cells": "",
                    "source_ranges": ", ".join(source_ranges),
                    "resolved": True,
                    "is_dynamic": False,
                }
            return {
                "dropdown_type": "Named Range",
                "options": "",
                "named_ranges": ref,
                "source_formula": ref,
                "referenced_cells": "",
                "source_ranges": "",
                "resolved": False,
                "is_dynamic": False,
            }

        structured = self._resolve_structured_reference(ref)
        if structured:
            return structured

        # Named range validation.
        # Direct sheet range reference, e.g. 'Sheet Name'!$A$1:$A$5
        sheet_name, coord = self._get_sheet_and_coord_from_ref(ref, current_sheet=current_sheet)
        if sheet_name and coord and sheet_name in self.wb.sheetnames:
            resolved_values = self._resolve_cell_range_values(self.wb[sheet_name], coord)
            if resolved_values:
                return {
                    "dropdown_type": "Direct Range",
                    "options": ", ".join(resolved_values),
                    "named_ranges": "",
                    "source_formula": ref,
                    "referenced_cells": "",
                    "source_ranges": f"{sheet_name}!{coord}",
                    "resolved": True,
                    "is_dynamic": False,
                }

        return {
            "dropdown_type": "Unresolved",
            "options": "",
            "named_ranges": "",
            "source_formula": ref,
            "referenced_cells": ", ".join(self._extract_referenced_cells(ref)),
            "source_ranges": "",
            "resolved": False,
            "is_dynamic": False,
        }

    def _build_named_range_map(self):
        """Maps named ranges to cell coordinates (e.g., 'File Info!I13' -> 'Z_FileInfo_Date')."""
        logging.info("Building Named Range index...")
        self.named_range_map.clear()
        self.named_range_cell_map.clear()

        for name, dn in self._iter_defined_names():
            try:
                # Ignore print areas and external names
                if getattr(dn, "type", "REGION") != "REGION" or getattr(dn, "is_external", False):
                    continue

                for sheet_name, coord in self._iter_defined_name_destinations(dn):
                    if not sheet_name or not coord:
                        continue

                    self.sheet_named_range_count[sheet_name] += 1

                    # For field-label lookup, exact single-cell destinations are most useful.
                    # If a range is defined, map the top-left cell and every covered cell.
                    top_left = coord
                    if ":" in coord:
                        try:
                            min_col, min_row, _, _ = range_boundaries(coord)
                            top_left = f"{get_column_letter(min_col)}{min_row}"
                            max_col, max_row = range_boundaries(coord)[2], range_boundaries(coord)[3]
                            for row in range(min_row, max_row + 1):
                                for col in range(min_col, max_col + 1):
                                    cell_key = f"{sheet_name}!{get_column_letter(col)}{row}"
                                    self.named_range_cell_map.setdefault(cell_key, name)
                        except Exception:
                            continue
                    else:
                        self.named_range_cell_map.setdefault(f"{sheet_name}!{top_left}", name)

                    key = f"{sheet_name}!{top_left}"
                    # Keep the first named range encountered for a cell
                    self.named_range_map.setdefault(key, name)

            except Exception:
                continue

        logging.info(f"Named range index built: {len(self.named_range_map)} cell mappings found.")

    def _build_table_map(self):
        """
        Cache workbook tables so structured references can be resolved.
        """
        self.table_map.clear()

        for ws in self.wb.worksheets:
            tables = getattr(ws, "tables", {})
            for table_name, table in tables.items():
                self.table_map[table_name] = {
                    "sheet": ws.title,
                    "ref": table.ref,
                    "columns": [getattr(col, "name", "") for col in getattr(table, "tableColumns", [])],
                }

    def get_named_range_for_cell(self, ws, row, col):
        coord_key = f"{ws.title}!{get_column_letter(col)}{row}"
        return self.named_range_cell_map.get(coord_key, "")

    def record_named_range_usage(self, named_range, sheet_name, csv_name):
        if not named_range:
            return
        self.named_range_usage[named_range]["tabs"].add(sheet_name)
        self.named_range_usage[named_range]["csvs"].add(csv_name)

    def _resolve_named_range_value_preview(self, dn, limit=10):
        """
        Return a compact preview of the values a named range points to.
        """
        preview_values = []

        try:
            for sheet_name, coord in self._iter_defined_name_destinations(dn):
                if sheet_name not in self.wb.sheetnames:
                    continue

                ws = self.wb[sheet_name]
                if ":" in coord:
                    for row in ws[coord]:
                        for cell in row:
                            if cell.value is not None:
                                preview_values.append(str(cell.value).strip())
                                if len(preview_values) >= limit:
                                    return preview_values
                else:
                    value = ws[coord].value
                    if value is not None:
                        preview_values.append(str(value).strip())
                        if len(preview_values) >= limit:
                            return preview_values
        except Exception:
            return preview_values

        return preview_values

    def _coord_span(self, coord):
        """
        Return a human readable start/end span and total cell count for one destination.
        """
        start_cell = ""
        end_cell = ""
        cell_count = 0

        try:
            if ":" in coord:
                min_col, min_row, max_col, max_row = range_boundaries(coord)
                start_cell = f"{get_column_letter(min_col)}{min_row}"
                end_cell = f"{get_column_letter(max_col)}{max_row}"
                cell_count = (max_col - min_col + 1) * (max_row - min_row + 1)
            else:
                start_cell = coord
                end_cell = coord
                cell_count = 1
        except Exception:
            pass

        return start_cell, end_cell, cell_count

    def infer_data_type(self, cell, is_dropdown=False):
        """
        Infer a human-readable field data type from cell value and formatting.
        """
        if is_dropdown:
            return "dropdown"

        value = cell.value
        if value is None:
            return ""

        if isinstance(value, str) and value.startswith("="):
            fmt = (cell.number_format or "").lower()
            if self._looks_like_currency_format(fmt):
                return "currency"
            if self._looks_like_date_format(fmt):
                return "date"
            if "%" in fmt:
                return "percentage"
            return "formula"

        if isinstance(value, bool):
            return "boolean"

        if getattr(cell, "is_date", False) or self._looks_like_date_format(cell.number_format):
            return "date"

        if isinstance(value, (int, float)) and not isinstance(value, bool):
            fmt = (cell.number_format or "").lower()
            if self._looks_like_currency_format(fmt):
                return "currency"
            if "%" in fmt:
                return "percentage"
            if isinstance(value, int) or (isinstance(value, float) and value.is_integer()):
                return "integer"
            return "decimal"

        if isinstance(value, (datetime, date)):
            return "date"

        if isinstance(value, str):
            return "text"

        return "text"

    def _looks_like_date_format(self, number_format):
        fmt = (number_format or "").lower()
        if not fmt:
            return False

        # Skip plain numeric formats that use "m" as minute/month in other contexts.
        return any(token in fmt for token in ["yy", "yyyy", "dd", "mmm", "hh", "mm/dd", "dd/mm", "m/d"])

    def _looks_like_currency_format(self, number_format):
        fmt = (number_format or "").lower()
        return any(symbol in fmt for symbol in ["$", "€", "£", "¥", "₹"])

    def infer_required_status(self, cell, is_formula=False, is_dropdown=False, dropdown_meta=None):
        """
        Classify a field as required, optional, calculated, or conditional.
        """
        if is_formula:
            formula = str(cell.value or "").upper()
            if any(token in formula for token in ["IF(", "IFS(", "SWITCH(", "CHOOSE(", "INDIRECT("]):
                return "conditional"
            return "calculated"

        if is_dropdown:
            if dropdown_meta and not dropdown_meta.get("resolved", True):
                return "conditional"
            return "required"

        if not cell.protection.locked:
            if cell.value is None or str(cell.value).strip() == "":
                return "optional"
            return "required"

        return "optional"

    def infer_relationship_type(self, formula):
        """
        Classify the relationship implied by a formula.
        """
        formula_str = str(formula)
        if self.is_pass_through(formula_str):
            return "Pass-Through"

        formula_upper = formula_str.upper()
        if any(token in formula_upper for token in ["XLOOKUP(", "VLOOKUP(", "INDEX(", "MATCH("]):
            return "Lookup"
        if any(token in formula_upper for token in ["SUM(", "AVERAGE(", "COUNT(", "COUNTA(", "MIN(", "MAX("]):
            return "Aggregation"
        if any(token in formula_upper for token in ["IF(", "IFS(", "SWITCH(", "CHOOSE(", "INDIRECT("]):
            return "Conditional"
        return "Formula Reference"

    def add_relationship(self, source_tab, target_tab, relationship_type):
        """
        Store a typed relationship between two tabs.
        """
        if not source_tab or not target_tab:
            return
        self.relationships[(source_tab, target_tab)].add(relationship_type or "Formula Reference")

    def find_field_label(self, ws, row, col):
        cache_key = f"{ws.title}!{row}!{col}"
        if cache_key in self.field_label_cache:
            return self.field_label_cache[cache_key]

        coord_key = f"{ws.title}!{get_column_letter(col)}{row}"
        if coord_key in self.named_range_map:
            label = self.named_range_map[coord_key]
            self.field_label_cache[cache_key] = label
            return label

        candidates = []

        # search left 15 columns
        for c in range(max(1, col - 15), col):
            val = self.get_effective_value(ws, row, c)
            if (
                isinstance(val, str)
                and val.strip()
                and not val.startswith("=")
                and len(val) < 120
            ):
                candidates.append((col - c, val.strip()))

        # search upward 8 rows
        for r in range(max(1, row - 8), row):
            val = self.get_effective_value(ws, r, col)
            if (
                isinstance(val, str)
                and val.strip()
                and not val.startswith("=")
                and len(val) < 120
            ):
                candidates.append((100 + (row - r), val.strip()))

        if candidates:
            candidates.sort()
            label = candidates[0][1]
        else:
            label = f"Unknown_Field_{get_column_letter(col)}{row}"

        self.field_label_cache[cache_key] = label
        return label

    def is_pass_through(self, formula):
        """
        Identifies direct-reference formulas that simply pass values through.
        Examples:
          =A1
          =$A$1
          ='Sheet Name'!A1
          =+A1
        """
        formula = str(formula).strip()

        patterns = [
            r"^=\+?\$?[A-Za-z]+\$?[0-9]+$",
            r"^=\+?'[^']+'!\$?[A-Za-z]+\$?[0-9]+$",
            r"^=\+?[A-Za-z0-9_ .\-]+!\$?[A-Za-z]+\$?[0-9]+$"
        ]

        return any(re.match(p, formula) for p in patterns)

    def classify_formula(self, formula):
        formula = str(formula).upper()

        if self.is_pass_through(formula):
            return "PassThrough"
        if "XLOOKUP(" in formula:
            return "Lookup"
        if "VLOOKUP(" in formula:
            return "Lookup"
        if "INDEX(" in formula:
            return "Lookup"
        if "MATCH(" in formula:
            return "Lookup"
        if "SUM(" in formula:
            return "Aggregation"
        if "AVERAGE(" in formula:
            return "Aggregation"
        if "IF(" in formula:
            return "Conditional"

        return "Calculation"

    def extract_named_range_dependencies(self, formula):
        """
        Case-insensitive search for named ranges used in a formula.
        Uses boundary-aware matching to reduce false positives.
        """
        formula_upper = str(formula).upper()
        found = []

        for nr in set(self.named_range_map.values()):
            nr_upper = str(nr).upper()
            pattern = rf"(?<![A-Z0-9_]){re.escape(nr_upper)}(?![A-Z0-9_])"
            if re.search(pattern, formula_upper):
                found.append(nr)

        return sorted(set(found))

    def extract_source_field_dependencies(self, ws, formula):
        """
        Resolve formula references into source field labels.

        This keeps the dependency output at the field level instead of sheet level
        for the Calculated Fields export.
        """
        formula_str = str(formula)
        deps = []
        seen = set()

        def add_dep(sheet_name, row_num, col_letter):
            if sheet_name not in self.wb.sheetnames:
                return

            source_ws = self.wb[sheet_name]
            col_idx = openpyxl.utils.column_index_from_string(col_letter)
            label = self.find_field_label(source_ws, row_num, col_idx)
            value = label if label else f"{sheet_name}!{col_letter}{row_num}"

            if value not in seen:
                seen.add(value)
                deps.append(value)

        # Sheet-qualified ranges and single-cell refs.
        sheet_ref_pattern = r"(?:'([^']+)'|([A-Za-z0-9_ .\-]+))!\$?([A-Za-z]+)\$?([0-9]+)(?::\$?([A-Za-z]+)\$?([0-9]+))?"
        for quoted_sheet, unquoted_sheet, start_col, start_row, end_col, end_row in re.findall(sheet_ref_pattern, formula_str):
            sheet_name = quoted_sheet if quoted_sheet else unquoted_sheet
            if not sheet_name:
                continue

            try:
                min_col, min_row = openpyxl.utils.column_index_from_string(start_col), int(start_row)
                if end_col and end_row:
                    max_col, max_row = openpyxl.utils.column_index_from_string(end_col), int(end_row)
                else:
                    max_col, max_row = min_col, min_row
            except Exception:
                continue

            for row_num in range(min_row, max_row + 1):
                for col_idx in range(min_col, max_col + 1):
                    add_dep(sheet_name.strip(), row_num, get_column_letter(col_idx))

        # Same-sheet cell refs, e.g. =A1 + B2
        same_sheet_pattern = r"(?<![A-Z0-9_!])\$?([A-Z]{1,3})\$?([0-9]+)(?![A-Z0-9_])"
        for col_letter, row_num in re.findall(same_sheet_pattern, formula_str):
            try:
                add_dep(ws.title, int(row_num), col_letter)
            except Exception:
                continue

        return deps

    def extract_field_mappings_from_formula(self, ws, formula):
        """
        Resolve formula references into (source tab, source field) pairs.

        This is used for Field Mapping.csv so the Source Field column contains
        actual labels instead of a placeholder.
        """
        formula_str = str(formula)
        mappings = []
        seen = set()

        def add_mapping(sheet_name, row_num, col_letter):
            if sheet_name not in self.wb.sheetnames:
                return

            source_ws = self.wb[sheet_name]
            col_idx = openpyxl.utils.column_index_from_string(col_letter)
            label = self.find_field_label(source_ws, row_num, col_idx)
            source_field = label if label else f"{sheet_name}!{col_letter}{row_num}"
            key = (sheet_name, source_field)
            if key not in seen:
                seen.add(key)
                mappings.append(key)

        sheet_ref_pattern = r"(?:'([^']+)'|([A-Za-z0-9_ .\-]+))!\$?([A-Za-z]+)\$?([0-9]+)(?::\$?([A-Za-z]+)\$?([0-9]+))?"
        for quoted_sheet, unquoted_sheet, start_col, start_row, end_col, end_row in re.findall(sheet_ref_pattern, formula_str):
            sheet_name = quoted_sheet if quoted_sheet else unquoted_sheet
            if not sheet_name:
                continue

            try:
                min_col, min_row = openpyxl.utils.column_index_from_string(start_col), int(start_row)
                if end_col and end_row:
                    max_col, max_row = openpyxl.utils.column_index_from_string(end_col), int(end_row)
                else:
                    max_col, max_row = min_col, min_row
            except Exception:
                continue

            for row_num in range(min_row, max_row + 1):
                for col_idx in range(min_col, max_col + 1):
                    add_mapping(sheet_name.strip(), row_num, get_column_letter(col_idx))

        same_sheet_pattern = r"(?<![A-Z0-9_!])\$?([A-Z]{1,3})\$?([0-9]+)(?![A-Z0-9_])"
        for col_letter, row_num in re.findall(same_sheet_pattern, formula_str):
            try:
                add_mapping(ws.title, int(row_num), col_letter)
            except Exception:
                continue

        return mappings

    def extract_dependencies_from_formula(self, formula):
        """
        Extract sheet-level dependencies from formulas.
        Handles quoted sheet names and sheet names with spaces / punctuation.
        """
        pattern = r"(?:'([^']+)'|([A-Za-z0-9_ .\-]+))!\$?[A-Za-z]+\$?[0-9]+"
        matches = re.findall(pattern, str(formula))

        deps = set()
        for quoted, unquoted in matches:
            sheet = quoted if quoted else unquoted
            if sheet:
                deps.add(sheet.strip())

        return sorted(deps)

    def get_effective_value(self, ws, row, col):
        cache_key = f"{ws.title}!{get_column_letter(col)}{row}"
        if cache_key in self.effective_value_cache:
            return self.effective_value_cache[cache_key]

        cell = ws.cell(row=row, column=col)
        if cell.value is not None:
            self.effective_value_cache[cache_key] = cell.value
            return cell.value

        if cache_key in self.merged_cell_lookup:
            value = self.merged_cell_lookup[cache_key]
            self.effective_value_cache[cache_key] = value
            return value

        self.effective_value_cache[cache_key] = None
        return None

    def export_named_ranges(self):
        logging.info("Exporting Named Ranges...")

        for name, dn in self._iter_defined_names():
            try:
                raw = getattr(dn, "value", None) or getattr(dn, "attr_text", None) or str(dn)
                raw = str(raw)
                values_preview = ", ".join(self._resolve_named_range_value_preview(dn))
                usage_tabs = ", ".join(sorted(self.named_range_usage.get(name, {}).get("tabs", set())))
                usage_csvs = ", ".join(sorted(self.named_range_usage.get(name, {}).get("csvs", set())))

                for sheet_name, coord in self._iter_defined_name_destinations(dn):
                    start_cell, end_cell, cell_count = self._coord_span(coord)
                    self._append_to_csv(
                        "Named Ranges.csv",
                        [
                            name,
                            sheet_name,
                            start_cell or coord,
                            end_cell or coord,
                            cell_count,
                            values_preview,
                            usage_tabs,
                            usage_csvs,
                            raw
                        ]
                    )

            except Exception:
                pass

    def analyze(self):
        try:
            self.load_workbook()

            for sheet_name in self.target_tabs:
                if sheet_name not in self.wb.sheetnames:
                    logging.warning(f"Tab '{sheet_name}' not found. Skipping.")
                    continue

                logging.info(f"Analyzing sheet: {sheet_name}")
                ws = self.wb[sheet_name]
                self.analyze_sheet(ws)
                logging.info(f"Finished saving data for: {sheet_name}")

            self.export_named_ranges()
            logging.info("Analysis complete! Generating final Tab Relationships...")
            self._write_relationships()
        finally:
            self._close_csv_handles()

    def analyze_sheet(self, ws):
        sheet_name = ws.title
        visibility = ws.sheet_state
        tab_color = ws.sheet_properties.tabColor.rgb if ws.sheet_properties.tabColor else "None"

        self._append_to_csv(
            "Tab Inventory.csv",
            [
                sheet_name,
                "Uncategorized",
                "Auto-extracted",
                "",
                "",
                visibility,
                tab_color
            ]
        )

        dropdowns = {}
        validations = getattr(ws.data_validations, "dataValidation", [])

        for dv in validations:
            if dv.type == "list":
                for cell_range in dv.sqref.ranges:
                    for row in range(cell_range.min_row, cell_range.max_row + 1):
                        for col in range(cell_range.min_col, cell_range.max_col + 1):
                            coord = f"{get_column_letter(col)}{row}"
                            dropdown_meta = self.resolve_dropdown_options(
                                dv.formula1,
                                current_sheet=sheet_name
                            )
                            dropdowns[coord] = dropdown_meta

        max_r, max_c = ws.max_row, ws.max_column
        if max_r > 5000:
            max_r = 5000

        for row_cells in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c):
            row = row_cells[0].row

            row_dim = ws.row_dimensions.get(row)
            is_hidden_row = bool(row_dim and row_dim.hidden)

            for cell in row_cells:
                col = cell.column
                col_letter = get_column_letter(col)

                col_dim = ws.column_dimensions.get(col_letter)
                is_hidden_col = bool(col_dim and col_dim.hidden)

                val = cell.value
                coord = cell.coordinate

                if val is None:
                    continue

                is_formula = isinstance(val, str) and val.startswith("=")
                is_dropdown = coord in dropdowns
                is_unlocked = not cell.protection.locked

                # Capture:
                #   - formulas
                #   - dropdowns
                #   - unlocked cells (text or numeric inputs)
                if is_formula or is_dropdown or is_unlocked:
                    label = self.find_field_label(ws, row, col)
                    notes = []
                    if is_hidden_row or is_hidden_col:
                        notes.append("[HIDDEN CELL]")
                    if is_formula and self.is_pass_through(val):
                        notes.append("[PASS-THROUGH]")

                    notes_str = " | ".join(notes)

                    named_range = self.get_named_range_for_cell(ws, row, col)
                    if named_range:
                        self.record_named_range_usage(named_range, sheet_name, "Field Inventory.csv")

                    formula_type = self.classify_formula(val) if is_formula else ""
                    required_status = self.infer_required_status(
                        cell,
                        is_formula=is_formula,
                        is_dropdown=is_dropdown,
                        dropdown_meta=dropdowns.get(coord)
                    )
                    inferred_data_type = self.infer_data_type(cell, is_dropdown=is_dropdown)

                    self._append_to_csv(
                        "Field Inventory.csv",
                        [
                            sheet_name,
                            label,
                            named_range,
                            "Yes" if not is_formula else "No",
                            "Yes" if is_formula else "No",
                            "Yes" if is_dropdown else "No",
                            formula_type,
                            inferred_data_type,
                            required_status,
                            notes_str,
                            coord
                        ]
                    )

                    if is_dropdown:
                        dropdown_meta = dropdowns[coord]
                        named_ranges = [
                            item.strip()
                            for item in str(dropdown_meta.get("named_ranges", "")).split(",")
                            if item.strip()
                        ]
                        for nr in named_ranges:
                            self.record_named_range_usage(nr, sheet_name, "Dropdown Values.csv")

                        self._append_to_csv(
                            "Dropdown Values.csv",
                            [
                                sheet_name,
                                label,
                                dropdown_meta.get("options", ""),
                                dropdown_meta.get("named_ranges", ""),
                                dropdown_meta.get("dropdown_type", ""),
                                dropdown_meta.get("source_formula", ""),
                                dropdown_meta.get("referenced_cells", ""),
                                dropdown_meta.get("source_ranges", ""),
                                "Resolved" if dropdown_meta.get("resolved") else "Unresolved",
                                coord
                            ]
                        )

                    if is_formula:
                        source_fields = self.extract_source_field_dependencies(ws, val)
                        deps = self.extract_dependencies_from_formula(val)
                        named_deps = self.extract_named_range_dependencies(val)
                        relationship_type = self.infer_relationship_type(val)
                        for named_dep in named_deps:
                            self.record_named_range_usage(named_dep, sheet_name, "Calculated Fields.csv")

                        self._append_to_csv(
                            "Calculated Fields.csv",
                            [
                                sheet_name,
                                label,
                                val,
                                ", ".join(source_fields),
                                ", ".join(named_deps),
                                coord
                            ]
                        )

                        # ShowHide usage capture
                        showhide_named_deps = [
                            dep for dep in named_deps
                            if "_SHOWHIDE" in dep.upper()
                        ]
                        if "_SHOWHIDE" in str(val).upper() or showhide_named_deps:
                            if showhide_named_deps:
                                for dep in showhide_named_deps:
                                    self.record_named_range_usage(dep, sheet_name, "ShowHide Usage.csv")
                                    self._append_to_csv(
                                        "ShowHide Usage.csv",
                                        [
                                            sheet_name,
                                            coord,
                                            val,
                                            dep
                                        ]
                                    )
                            else:
                                self._append_to_csv(
                                    "ShowHide Usage.csv",
                                    [
                                        sheet_name,
                                        coord,
                                        val,
                                        ""
                                    ]
                                )

                        for dep in deps:
                            if dep != sheet_name:
                                self.add_relationship(dep, sheet_name, relationship_type)
                        for source_tab, source_field in self.extract_field_mappings_from_formula(ws, val):
                            if source_tab != sheet_name:
                                self.add_relationship(source_tab, sheet_name, relationship_type)
                            self._append_to_csv(
                                "Field Mapping.csv",
                                [
                                    source_tab,
                                    source_field,
                                    sheet_name,
                                    ""
                                ]
                            )

    def _write_relationships(self):
        for source, target in sorted(self.relationships):
            relationship_types = ", ".join(sorted(self.relationships[(source, target)]))
            self._append_to_csv(
                "Tab Relationships.csv",
                [
                    source,
                    target,
                    relationship_types,
                    f"'{target}' pulls data from '{source}'"
                ]
            )


if __name__ == "__main__":
    analyzer = WorkbenchAnalyzer(FILE_PATH, TARGET_TABS, OUTPUT_DIR)
    analyzer.analyze()
