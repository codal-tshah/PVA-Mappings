"""
PVA Workbench Analyzer — Output Writer
Populates the documentation template workbook with all extracted data.
"""
from __future__ import annotations
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import COLORS, OUTPUT_DIR
from logger import get_logger

log = get_logger("output_writer")

# ── Style helpers ──────────────────────────────────────────────────────────────

def _fill(hex_: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_)

def _bd() -> Border:
    s = Side(style="thin", color=COLORS["border"])
    return Border(left=s, right=s, top=s, bottom=s)

def _ft(bold=False, color="1E293B", sz=9, italic=False) -> Font:
    return Font(bold=bold, color=color, size=sz, name="Arial", italic=italic)

def _al(h="left", wrap=True) -> Alignment:
    return Alignment(horizontal=h, vertical="center", wrap_text=wrap)

def _hcell(ws, r: int, c: int, v: Any, bg=None, fg=None, sz=9, bold=True, center=True):
    bg = bg or COLORS["header_bg"]
    fg = fg or COLORS["white"]
    cell = ws.cell(row=r, column=c, value=v)
    cell.fill = _fill(bg)
    cell.font = _ft(bold=bold, color=fg, sz=sz)
    cell.alignment = _al("center" if center else "left")
    cell.border = _bd()
    return cell

def _dcell(ws, r: int, c: int, v: Any, bg=None, bold=False, h="left"):
    bg = bg or COLORS["white"]
    cell = ws.cell(row=r, column=c, value=v)
    cell.fill = _fill(bg)
    cell.font = _ft(bold=bold)
    cell.alignment = _al(h)
    cell.border = _bd()
    return cell

def _title_row(ws, r: int, text: str, ncols: int, bg=None):
    bg = bg or COLORS["header_bg"]
    ws.row_dimensions[r].height = 28
    ws.merge_cells(f"A{r}:{get_column_letter(ncols)}{r}")
    c = ws[f"A{r}"]
    c.value = text
    c.font = Font(bold=True, color=COLORS["white"], size=12, name="Arial")
    c.fill = _fill(bg)
    c.alignment = _al("center")

def _section_row(ws, r: int, text: str, ncols: int, bg=None):
    bg = bg or COLORS["section_bg"]
    ws.row_dimensions[r].height = 18
    ws.merge_cells(f"A{r}:{get_column_letter(ncols)}{r}")
    c = ws[f"A{r}"]
    c.value = f"  ▸  {text}"
    c.font = _ft(bold=True, color=COLORS["white"])
    c.fill = _fill(bg)
    c.alignment = _al("left")
    c.border = _bd()

def _type_bg(field_type: str) -> str:
    return {
        "input":      COLORS["input_bg"],
        "calculated": COLORS["calc_bg"],
        "dropdown":   COLORS["dropdown_bg"],
    }.get(field_type, COLORS["white"])

def _req_bg(req: str) -> str:
    return {
        "Required":    COLORS["req_bg"],
        "Conditional": COLORS["cond_bg"],
        "Optional":    COLORS["opt_bg"],
        "Calculated":  COLORS["calc_bg"],
    }.get(req, COLORS["white"])


# ── Sheet writers ──────────────────────────────────────────────────────────────

class OutputWriter:
    def __init__(self, template_path: str | None, output_path: str):
        self.output_path = output_path
        if template_path and Path(template_path).exists():
            self.wb = load_workbook(template_path)
            log.info(f"Loaded template: {template_path}")
        else:
            self.wb = Workbook()
            self.wb.remove(self.wb.active)
            log.info("No template found — creating fresh workbook")

        self._ensure_sheets()

    def _ensure_sheets(self):
        required = [
            "Tab Inventory", "Field Inventory", "Field Mapping",
            "Dropdown Values", "Calculated Fields", "Tab Relationships",
            "Notes & Findings", "Named Ranges",
        ]
        for name in required:
            if name not in self.wb.sheetnames:
                self.wb.create_sheet(name)

    # ── Tab Inventory ──────────────────────────────────────────────────────────
    def init_tab_inventory(self):
        ws = self.wb["Tab Inventory"]
        ws.sheet_view.showGridLines = False
        ws.delete_rows(1, ws.max_row + 1)

        cols = {"A":28,"B":20,"C":12,"D":12,"E":16,"F":7,"G":7,"H":7,
                "I":7,"J":7,"K":8,"L":8,"M":8,"N":35,"O":40}
        for col, w in cols.items():
            ws.column_dimensions[col].width = w

        _title_row(ws, 1, "Tab Inventory — PVA Workbench (28 Target Tabs)", 15)
        headers = [
            "Tab Name","Matched Sheet Name","Visible?","Tab Color","Used Range",
            "Total Cells","Formula Cells","Merged Cells","Hidden Rows","Hidden Cols",
            "Tables","DV Count","Named Ranges Referenced","Notes",
        ]
        for c, h in enumerate(headers, 1):
            _hcell(ws, 2, c, h, sz=8)
        ws.freeze_panes = "A3"

    def write_tab_inventory_row(self, rec):
        ws = self.wb["Tab Inventory"]
        r = ws.max_row + 1
        bg = COLORS["alt_bg"] if r % 2 == 0 else COLORS["white"]
        vis_bg = {"Visible": COLORS["opt_bg"], "Hidden": COLORS["cond_bg"],
                  "VeryHidden": COLORS["req_bg"]}.get(rec.visible, COLORS["white"])

        _dcell(ws, r, 1, rec.tab, bg, bold=True)
        _dcell(ws, r, 2, rec.matched_name, bg)
        c3 = ws.cell(row=r, column=3, value=rec.visible)
        c3.fill = _fill(vis_bg); c3.font = _ft(bold=True, sz=9)
        c3.alignment = _al("center"); c3.border = _bd()
        _dcell(ws, r, 4, rec.tab_color or "", bg)
        _dcell(ws, r, 5, rec.used_range, bg)
        for i, v in enumerate([rec.total_cells, rec.formula_cells, rec.merged_cells,
                                rec.hidden_rows, rec.hidden_cols], 6):
            _dcell(ws, r, i, v, bg, h="center")
        _dcell(ws, r, 11, rec.tables, bg)
        _dcell(ws, r, 12, rec.data_validations, bg, h="center")
        _dcell(ws, r, 13, rec.named_ranges_referenced, COLORS["alt_bg"])
        _dcell(ws, r, 14, rec.notes, bg)
        ws.row_dimensions[r].height = 16

    # ── Field Inventory ────────────────────────────────────────────────────────
    def init_field_inventory(self):
        ws = self.wb["Field Inventory"]
        ws.sheet_view.showGridLines = False
        ws.delete_rows(1, ws.max_row + 1)

        cols = {"A":24,"B":38,"C":7,"D":7,"E":7,"F":8,"G":12,"H":10,"I":26,"J":7,"K":42}
        for col, w in cols.items():
            ws.column_dimensions[col].width = w

        _title_row(ws, 1, "Field Inventory — All Fields Across 28 Target Tabs", 11)
        headers = ["Tab","Field / Label","Input","Calc","DD",
                   "Data Type","Required?","Cell Ref","Named Ranges Used","API Seed?","Notes"]
        for c, h in enumerate(headers, 1):
            _hcell(ws, 2, c, h, sz=8)
        ws.freeze_panes = "A3"
        self._field_inv_last_tab: str = ""

    def write_field_inventory_row(self, rec):
        ws = self.wb["Field Inventory"]

        # Section header when tab changes
        if rec.tab != self._field_inv_last_tab:
            r_sec = ws.max_row + 1
            _section_row(ws, r_sec, f"TAB: {rec.tab}", 11)
            ws.row_dimensions[r_sec].height = 18
            self._field_inv_last_tab = rec.tab

        r = ws.max_row + 1
        bg = _type_bg(rec.field_type)
        rb = _req_bg(rec.is_required)
        ws.row_dimensions[r].height = 14

        _dcell(ws, r, 1, rec.tab, COLORS["alt_bg"] if r % 2 == 0 else COLORS["white"], bold=True)
        _dcell(ws, r, 2, rec.label, bg)
        for c, v in [(3, "Y" if rec.field_type=="input" else ""),
                     (4, "Y" if rec.field_type=="calculated" else ""),
                     (5, "Y" if rec.field_type=="dropdown" else "")]:
            cc = ws.cell(row=r, column=c, value=v)
            cc.fill = _fill("A7F3D0" if v=="Y" else COLORS["white"])
            cc.font = _ft(bold=(v=="Y"))
            cc.alignment = _al("center"); cc.border = _bd()
        _dcell(ws, r, 6, rec.data_type, bg)
        req_c = ws.cell(row=r, column=7, value=rec.is_required)
        req_c.fill = _fill(rb); req_c.font = _ft(bold=(rec.is_required=="Required"), sz=9)
        req_c.alignment = _al("center"); req_c.border = _bd()
        _dcell(ws, r, 8, rec.cell_ref, COLORS["alt_bg"])
        _dcell(ws, r, 9, ", ".join(rec.named_ranges_used[:5]), bg)
        api_v = "N" if rec.field_type == "calculated" else "Y"
        api_c = ws.cell(row=r, column=10, value=api_v)
        api_c.fill = _fill("BBF7D0" if api_v=="Y" else "FCE7F3")
        api_c.font = _ft(bold=True); api_c.alignment = _al("center"); api_c.border = _bd()
        _dcell(ws, r, 11, rec.notes, bg)

    # ── Dropdown Values ────────────────────────────────────────────────────────
    def init_dropdown_values(self):
        ws = self.wb["Dropdown Values"]
        ws.sheet_view.showGridLines = False
        ws.delete_rows(1, ws.max_row + 1)

        cols = {"A":24,"B":7,"C":35,"D":12,"E":80,"F":28}
        for col, w in cols.items():
            ws.column_dimensions[col].width = w

        _title_row(ws, 1, "Dropdown Values — All Data Validation Fields & Options", 6)
        for c, h in enumerate(["Tab","Cell Ref","Field Label","DV Type","All Options (pipe-separated)","Source / Notes"], 1):
            _hcell(ws, 2, c, h, sz=8)
        ws.freeze_panes = "A3"
        self._dd_last_tab: str = ""

    def write_dropdown_row(self, rec):
        ws = self.wb["Dropdown Values"]
        if rec.tab != self._dd_last_tab:
            r_sec = ws.max_row + 1
            _section_row(ws, r_sec, f"TAB: {rec.tab}", 6)
            self._dd_last_tab = rec.tab

        r = ws.max_row + 1
        bg = COLORS["dropdown_bg"] if r % 2 == 0 else COLORS["white"]
        ws.row_dimensions[r].height = 16
        _dcell(ws, r, 1, rec.tab, bg, bold=True)
        _dcell(ws, r, 2, rec.cell_ref, bg)
        _dcell(ws, r, 3, rec.field_label, bg)
        _dcell(ws, r, 4, rec.validation_type, bg, h="center")
        _dcell(ws, r, 5, rec.options_str, bg)
        note = rec.source_range or rec.notes
        _dcell(ws, r, 6, note, bg)

    # ── Calculated Fields ──────────────────────────────────────────────────────
    def init_calculated_fields(self):
        ws = self.wb["Calculated Fields"]
        ws.sheet_view.showGridLines = False
        ws.delete_rows(1, ws.max_row + 1)

        cols = {"A":24,"B":7,"C":32,"D":12,"E":70,"F":28,"G":30}
        for col, w in cols.items():
            ws.column_dimensions[col].width = w

        _title_row(ws, 1, "Calculated Fields — Formulas, UDFs & Named Range Logic", 7)
        for c, h in enumerate(["Tab","Cell Ref","Field / Label","Category","Formula","UDFs / Named Ranges","Cross-Sheet Refs"], 1):
            _hcell(ws, 2, c, h, sz=8)
        ws.freeze_panes = "A3"
        self._calc_last_tab: str = ""

    def write_formula_row(self, rec):
        ws = self.wb["Calculated Fields"]
        if rec.tab != self._calc_last_tab:
            r_sec = ws.max_row + 1
            _section_row(ws, r_sec, f"TAB: {rec.tab}", 7)
            self._calc_last_tab = rec.tab

        r = ws.max_row + 1
        bg = COLORS["calc_bg"] if r % 2 == 0 else COLORS["white"]
        ws.row_dimensions[r].height = 16
        _dcell(ws, r, 1, rec.tab, bg, bold=True)
        _dcell(ws, r, 2, rec.cell_ref, COLORS["alt_bg"])
        _dcell(ws, r, 3, rec.label, bg)
        _dcell(ws, r, 4, rec.formula_category, bg, h="center")
        f_cell = ws.cell(row=r, column=5, value=rec.formula[:300])
        f_cell.fill = _fill(COLORS["white"])
        f_cell.font = Font(name="Consolas", size=8, color="1C2833")
        f_cell.alignment = _al("left"); f_cell.border = _bd()
        udf_str = ", ".join(rec.udfs_used + rec.named_ranges_used[:4])
        _dcell(ws, r, 6, udf_str, bg)
        _dcell(ws, r, 7, ", ".join(rec.cross_sheet_refs[:6]), bg)

    # ── Field Mapping (relationships) ─────────────────────────────────────────
    def init_field_mapping(self):
        ws = self.wb["Field Mapping"]
        ws.sheet_view.showGridLines = False
        ws.delete_rows(1, ws.max_row + 1)

        cols = {"A":28,"B":35,"C":30,"D":18,"E":60,"F":14}
        for col, w in cols.items():
            ws.column_dimensions[col].width = w

        _title_row(ws, 1, "Field Mapping — Cross-Tab Data Flow & Dependencies", 6)
        for c, h in enumerate(["Source Tab","Source Field","Target Tab","Relationship Type","What Flows / How","Confidence"], 1):
            _hcell(ws, 2, c, h, sz=8)
        ws.freeze_panes = "A3"

    def write_relationship_row(self, rec):
        ws = self.wb["Field Mapping"]
        r = ws.max_row + 1
        bg = COLORS["alt_bg"] if r % 2 == 0 else COLORS["white"]
        ws.row_dimensions[r].height = 28

        REL_COLORS = {
            "Drives":         "FDE8D0",
            "Feeds":          "D5F5E3",
            "Auto-populates": "D6EAF8",
            "Configures":     COLORS["calc_bg"],
            "Linked from":    "EDE7F6",
        }
        rel_bg = REL_COLORS.get(rec.relationship_type, COLORS["white"])
        conf_bg = {"confirmed": "D5F5E3", "inferred": COLORS["calc_bg"],
                   "known-pattern": "D6EAF8"}.get(rec.confidence, COLORS["white"])

        _dcell(ws, r, 1, rec.source_tab, bg, bold=True)
        _dcell(ws, r, 2, rec.source_field, COLORS["input_bg"])
        _dcell(ws, r, 3, rec.target_tab, COLORS["calc_bg"])
        rc = ws.cell(row=r, column=4, value=rec.relationship_type)
        rc.fill = _fill(rel_bg); rc.font = _ft(bold=True); rc.alignment = _al("center"); rc.border = _bd()
        _dcell(ws, r, 5, rec.what_flows, bg)
        cc = ws.cell(row=r, column=6, value=rec.confidence)
        cc.fill = _fill(conf_bg); cc.font = _ft(sz=8, italic=True)
        cc.alignment = _al("center"); cc.border = _bd()

    # ── Tab Relationships ──────────────────────────────────────────────────────
    def init_tab_relationships(self):
        ws = self.wb["Tab Relationships"]
        ws.sheet_view.showGridLines = False
        ws.delete_rows(1, ws.max_row + 1)

        cols = {"A":28,"B":30,"C":18,"D":60}
        for col, w in cols.items():
            ws.column_dimensions[col].width = w

        _title_row(ws, 1, "Tab Relationships — Inter-Tab Dependencies", 4)
        for c, h in enumerate(["Source Tab","Target Tab","Type","What Flows"], 1):
            _hcell(ws, 2, c, h, sz=8)
        ws.freeze_panes = "A3"

    def write_tab_relationship_row(self, rec):
        ws = self.wb["Tab Relationships"]
        r = ws.max_row + 1
        bg = COLORS["alt_bg"] if r % 2 == 0 else COLORS["white"]
        ws.row_dimensions[r].height = 22
        _dcell(ws, r, 1, rec.source_tab, bg, bold=True)
        _dcell(ws, r, 2, rec.target_tab, COLORS["calc_bg"])
        _dcell(ws, r, 3, rec.relationship_type, bg, h="center")
        _dcell(ws, r, 4, rec.what_flows, bg)

    # ── Notes & Findings ──────────────────────────────────────────────────────
    def init_notes(self):
        ws = self.wb["Notes & Findings"]
        ws.sheet_view.showGridLines = False
        ws.delete_rows(1, ws.max_row + 1)

        cols = {"A":24,"B":55,"C":45}
        for col, w in cols.items():
            ws.column_dimensions[col].width = w

        _title_row(ws, 1, "Notes & Findings — Gaps, Warnings & Recommendations", 3)
        for c, h in enumerate(["Area","Finding / Detail","Recommendation / Action"], 1):
            _hcell(ws, 2, c, h, sz=8)
        ws.freeze_panes = "A3"

    def write_note(self, area: str, finding: str, recommendation: str = ""):
        ws = self.wb["Notes & Findings"]
        r = ws.max_row + 1
        bg = COLORS["gap_bg"] if r % 2 == 0 else COLORS["white"]
        ws.row_dimensions[r].height = 40
        _dcell(ws, r, 1, area, bg, bold=True)
        _dcell(ws, r, 2, finding, bg)
        _dcell(ws, r, 3, recommendation, COLORS["gap_bg"] if recommendation else COLORS["white"])

    # ── Named Ranges ──────────────────────────────────────────────────────────
    def init_named_ranges(self):
        ws = self.wb["Named Ranges"]
        ws.sheet_view.showGridLines = False
        ws.delete_rows(1, ws.max_row + 1)

        cols = {"A":38,"B":12,"C":55,"D":40}
        for col, w in cols.items():
            ws.column_dimensions[col].width = w

        _title_row(ws, 1, "Named Ranges — pva*, N1*, tbl*, LIST_* Inventory", 4)
        for c, h in enumerate(["Named Range","Prefix","Refers To","Purpose / What It Controls"], 1):
            _hcell(ws, 2, c, h, sz=8)
        ws.freeze_panes = "A3"

    def write_named_range_row(self, rec):
        ws = self.wb["Named Ranges"]
        r = ws.max_row + 1
        PFX_COLORS = {
            "pva":      COLORS["dropdown_bg"],
            "N1":       COLORS["input_bg"],
            "tbl":      COLORS["alt_bg"],
            "LIST_":    COLORS["calc_bg"],
            "data_":    COLORS["calc_bg"],
            "lucro_":   "F5EEF8",
            "pvaMktSur":"D5F5E3",
        }
        bg = PFX_COLORS.get(rec.prefix, COLORS["white"])
        ws.row_dimensions[r].height = 16
        _dcell(ws, r, 1, rec.name, bg, bold=True)
        pc = ws.cell(row=r, column=2, value=rec.prefix)
        pc.fill = _fill(bg); pc.font = _ft(bold=True, sz=8)
        pc.alignment = _al("center"); pc.border = _bd()
        r_cell = ws.cell(row=r, column=3, value=rec.refers_to[:100])
        r_cell.fill = _fill(COLORS["white"])
        r_cell.font = Font(name="Consolas", size=8, color="1C2833")
        r_cell.alignment = _al("left"); r_cell.border = _bd()
        _dcell(ws, r, 4, rec.purpose, bg)

    # ── Save ──────────────────────────────────────────────────────────────────
    def save(self):
        self.wb.save(self.output_path)
        size = Path(self.output_path).stat().st_size / 1024
        log.info(f"Saved: {self.output_path} ({size:.0f} KB)")

    def save_intermediate(self, suffix: str):
        path = str(self.output_path).replace(".xlsx", f"_{suffix}.xlsx")
        self.wb.save(path)
        log.debug(f"Intermediate save: {path}")
